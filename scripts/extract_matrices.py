"""Extrai as matrizes (em branco) das planilhas originais para prisma/data/*.json.

Uso: python3 scripts/extract_matrices.py <Matriz_Wise.xlsx> <Check_list_12_basicos.xlsx>
"""
import json, re, sys
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent.parent / "prisma" / "data"
DIMS = [("B", "ATIVIDADE"), ("D", "QUALIDADE"), ("F", "IMPACTO")]


def clean(s):
    return re.sub(r"\s+", " ", str(s).replace("​", "")).strip()


ANSWER_CUTS = [
    r"\s*\((?:Sim|sim|Exemplo de|evidenciar|forms|politics day|Politic Day)\b.*$",
    r"\s*(?:Sim\s*,|Sim\s|R\.:).*$",
    r"\s*[*_]?(?:avaliação )?in loco.*$",
    r"\s*\*IN LOCO.*$",
    r"\s*\(Há uma RACI disponível\)",
]


def strip_answer(text):
    """Remove respostas/comentários do site anterior, mantendo apenas a afirmação da matriz."""
    t = clean(text)
    for pat in ANSWER_CUTS:
        t = re.sub(pat, "", t).strip()
    t = re.sub(r"\.(\s*\.)+$", ".", t)
    return t


# Parênteses originais da matriz DuPont que devem ser mantidos no fim das afirmações.
LEGIT_PAREN = re.compile(
    r"etc|%|^FR1$|^LTA e Não-LTA$|^investimentos simples visíveis$|^Específico, Mensurável|"
    r"^para o mesmo risco$|^incluindo (requisitos de EPI|contratadas)$|^mínimo semanalmente$|"
    r"^6 meses no máximo$|^prioridades, planejamento|^proteção contra incêndio|^orçamento, multiplicadores|"
    r"^experiência e diferentes níveis$|^testes, certificações|^exposições, visitas|^com a pessoa lesionada|"
    r"^integração$|^ou aprendizados/tendências|^em período integral"
)
TAIL_PAREN = re.compile(r"^(.*?)\s*\(((?:[^()]|\([^()]*\))*)\)\s*\.?\s*$")


def strip_site_note(text):
    """Em afirmações já avaliadas, remove as notas do site entre parênteses no fim do texto."""
    while True:
        m = TAIL_PAREN.match(text)
        if not m or LEGIT_PAREN.search(m.group(2).strip()):
            return text
        text = m.group(1).rstrip(" .") + ("." if text.rstrip().endswith(".") else "")


def extract_wise(path):
    wb_f = openpyxl.load_workbook(path)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    elements = []
    for ws in wb_f.worksheets:
        m = re.match(r"^(\d+)\.", ws.title)
        if not m:
            continue
        num = int(m.group(1))
        wv = wb_v[ws.title]
        title = clean(wv["A1"].value)
        title = re.sub(r"^\d+\.\s*", "", title)
        levels = []
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=ROUND(SUM("):
                    mm = re.search(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)/(\d+)", v)
                    r0, r1, den = int(mm.group(2)), int(mm.group(4)), int(mm.group(5))
                    lvl = wv[f"A{r0}"].value
                    stmts = []
                    for col, dim in DIMS:
                        for r in range(r0, r1 + 1):
                            t = wv[f"{col}{r}"].value
                            if t and isinstance(t, str) and len(clean(t)) > 3:
                                score_col = chr(ord(col) + 1)
                                scored = wv[f"{score_col}{r}"].value not in (None, "-")
                                txt = strip_answer(t)
                                if scored:
                                    txt = strip_site_note(txt)
                                stmts.append({"dimension": dim, "row": r, "text": txt})
                    stmts.sort(key=lambda s: (s["row"], ["ATIVIDADE", "QUALIDADE", "IMPACTO"].index(s["dimension"])))
                    if len(stmts) * 2 != den:
                        print(f"  aviso: {ws.title} nível {lvl}: {len(stmts)} afirmações, denominador {den}")
                    levels.append({"level": int(lvl), "maxPoints": den, "statements": stmts})
        levels.sort(key=lambda l: l["level"])
        # glossário do elemento
        glossary = None
        for row in wv.iter_rows(values_only=True):
            pass
        for r in range(1, wv.max_row + 1):
            if clean(wv[f"A{r}"].value or "") == "Glossário":
                g = wv[f"A{r+1}"].value
                glossary = clean(g) if g else None
                break
        reqs = []
        for l in levels:
            for i, s in enumerate(l["statements"], 1):
                reqs.append({
                    "code": f"E{num}.N{l['level']}.{i}",
                    "level": l["level"],
                    "dimension": s["dimension"],
                    "text": s["text"],
                })
        elements.append({
            "number": num,
            "title": title,
            "glossary": glossary,
            "levelMax": {str(l["level"]): l["maxPoints"] for l in levels},
            "requirements": reqs,
        })
        print(f"WISE E{num} {title}: {len(reqs)} afirmações")
    elements.sort(key=lambda e: e["number"])
    return elements


def extract_basics(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["SB Checklist Distribution"]
    basics, cur = [], None
    for r in ws.iter_rows(min_row=10, values_only=True):
        a, text = r[0], r[9]
        if isinstance(a, str) and re.match(r"^\s*\d+\s*\.", a) and not text:
            n = int(re.match(r"^\s*(\d+)", a).group(1))
            if cur is None or cur["number"] != n:
                cur = {"number": n, "header": clean(a), "requirements": []}
                basics.append(cur)
            continue
        if cur and isinstance(text, str) and len(clean(text)) > 20:
            risk = r[10] if r[10] in (1, 2, 3) else None
            cur["requirements"].append({"text": clean(text), "risk": risk})
    # critérios em português (Planilha1), casados por ordem dentro de cada básico
    wc = wb["Planilha1"]
    crit, sec = {}, None
    for r in wc.iter_rows(min_row=7, values_only=True):
        a = r[0]
        if isinstance(a, str) and re.match(r"^\s*\d+\s*\.", a) and not r[2]:
            sec = int(re.match(r"^\s*(\d+)", a).group(1))
            crit.setdefault(sec, [])
            continue
        if sec and isinstance(r[2], str) and len(r[2]) > 20:
            crit[sec].append({
                "basic": clean(r[5]) if r[5] else None,
                "partial": clean(r[6]) if r[6] else None,
                "significant": clean(r[7]) if r[7] else None,
            })
    for b in basics:
        cs = crit.get(b["number"], [])
        for i, req in enumerate(b["requirements"]):
            req["code"] = f"{b['number']}.{i+1}"
            if i < len(cs):
                req["criteria"] = cs[i]
            if req["risk"] is None:
                req["risk"] = 2
        print(f"Básico {b['number']}: {len(b['requirements'])} requisitos (critérios: {len(cs)})")
    return basics


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    wise = extract_wise(sys.argv[1])
    (OUT / "wise-matrix.json").write_text(json.dumps(wise, ensure_ascii=False, indent=1))
    basics = extract_basics(sys.argv[2])
    (OUT / "basics-checklist.json").write_text(json.dumps(basics, ensure_ascii=False, indent=1))
